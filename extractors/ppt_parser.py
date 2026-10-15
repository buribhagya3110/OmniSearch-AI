import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Dict, Any, List
from utils.logger import setup_logger

logger = setup_logger("PPTParser")

class PPTParser:
    """Extracts text, slides, tables, and notes from PPTX presentation files."""

    @classmethod
    def parse(cls, file_path: str) -> Dict[str, Any]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PPTX file not found: {file_path}")

        # Attempt 1: python-pptx
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            slides_text = []
            notes_text = []

            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            line = "".join([run.text for run in paragraph.runs]).strip()
                            if line:
                                slide_parts.append(line)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_vals:
                                slide_parts.append(" | ".join(row_vals))

                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    note = slide.notes_slide.notes_text_frame.text.strip()
                    if note:
                        notes_text.append(f"Slide {slide_num} Notes: {note}")

                slide_content = "\n".join(slide_parts).strip()
                if slide_content:
                    slides_text.append(f"--- Slide {slide_num} ---\n{slide_content}")

            full_text = "\n\n".join(slides_text + notes_text).strip()
            return {
                "text": full_text,
                "slides": slides_text,
                "notes": notes_text,
                "total_slides": len(prs.slides),
                "metadata": {
                    "total_slides": len(prs.slides),
                    "notes_count": len(notes_text)
                }
            }
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"python-pptx failed for {file_path}: {e}. Trying fallback.")

        # Attempt 2: Standard library zipfile + XML ElementTree fallback
        try:
            slides_text = []
            with zipfile.ZipFile(str(path)) as z:
                slide_files = sorted([f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")])
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for i, s_file in enumerate(slide_files, start=1):
                    xml_data = z.read(s_file)
                    root = ET.fromstring(xml_data)
                    texts = [node.text for node in root.findall(".//a:t", ns) if node.text]
                    if texts:
                        slides_text.append(f"--- Slide {i} ---\n" + "\n".join(texts))

            full_text = "\n\n".join(slides_text).strip()
            return {
                "text": full_text,
                "slides": slides_text,
                "notes": [],
                "total_slides": len(slides_text),
                "metadata": {"fallback": True, "total_slides": len(slides_text)}
            }
        except Exception as e:
            logger.error(f"Fallback PPTX parsing failed for {file_path}: {e}")
            return {
                "text": "",
                "slides": [],
                "notes": [],
                "total_slides": 0,
                "metadata": {"error": str(e)}
            }
