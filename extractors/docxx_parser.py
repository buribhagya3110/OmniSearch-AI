import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Dict, Any, List
from utils.logger import setup_logger

logger = setup_logger("DOCXParser")

class DOCXParser:
    """Extracts text, paragraphs, tables, and metadata from DOCX files."""

    @classmethod
    def parse(cls, file_path: str) -> Dict[str, Any]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        # Attempt 1: python-docx
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extract tables
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        table_texts.append(row_text)

            full_text = "\n\n".join(paragraphs + table_texts).strip()

            core_props = doc.core_properties
            metadata = {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "created": str(core_props.created) if core_props.created else "",
                "paragraphs_count": len(paragraphs),
                "tables_count": len(doc.tables)
            }

            return {
                "text": full_text,
                "paragraphs": paragraphs,
                "tables": table_texts,
                "metadata": metadata
            }
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"python-docx failed for {file_path}: {e}. Trying fallback.")

        # Attempt 2: Standard library zipfile + XML ElementTree fallback
        try:
            paragraphs = []
            with zipfile.ZipFile(str(path)) as z:
                # Word document text resides in word/document.xml
                if "word/document.xml" in z.namelist():
                    xml_content = z.read("word/document.xml")
                    root = ET.fromstring(xml_content)
                    # Namespace for WordprocessingML
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    for p in root.findall(".//w:p", ns):
                        texts = [node.text for node in p.findall(".//w:t", ns) if node.text]
                        if texts:
                            paragraphs.append("".join(texts))

            full_text = "\n\n".join(paragraphs).strip()
            return {
                "text": full_text,
                "paragraphs": paragraphs,
                "tables": [],
                "metadata": {"fallback": True, "paragraphs_count": len(paragraphs)}
            }
        except Exception as e:
            logger.error(f"Fallback DOCX parsing failed for {file_path}: {e}")
            return {
                "text": "",
                "paragraphs": [],
                "tables": [],
                "metadata": {"error": str(e)}
            }
